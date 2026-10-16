"""
Bhasha Setu shared model and data utilities for the supported FastAPI/PWA pilot.

Curriculum translation is currently Hindi-to-Santali only. Ho and Mundari have
small word lookup lists, not translation models.

Put these next to this file:
    index.html          the web page
    lora_santali/       the adapter you trained (adapter_config.json + adapter_model.safetensors)
    vocab/              ho_hindi_pairs_raw.csv and mundari_pairs_raw.csv (word lists)

Run:
    python api.py              start the FastAPI server and current PWA
    python api.py --mock       run the interface with clearly labelled fake drafts
    python app.py --selftest   load the model, translate three sentences, and exit

Then open http://localhost:8000

How each language works:
    Santali   the model drafts in the sat_Beng Bengali-script setting; a teacher and a community
              language expert must review each item before it reaches students.
    Ho        no translation model exists; the app has a small Ho-Hindi word lookup list.
    Mundari   no translation model exists; the app has a small Mundari-Hindi word lookup list.
"""
import csv
import difflib
import gc
import importlib
import importlib.util
import io
import json
import os
import queue
import re
import shutil
import sqlite3
import sys
import threading
import time
import traceback
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

HERE = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("BS_HOME", HERE))
DB_PATH = DATA_DIR / "bhasha_setu_v2.db"
UPLOAD_DIR = DATA_DIR / "uploads"
ADAPTER_DIR = HERE / "lora_santali"
VOCAB_DIR = HERE / "vocab"
BASE_MODEL = "facebook/nllb-200-distilled-600M"
PORT = int(os.environ.get("BS_PORT", "8000"))
MOCK = "--mock" in sys.argv or os.environ.get("BS_MOCK") == "1"

MAX_UPLOAD = 50_000_000      # bytes
MAX_SENTENCES = 20_000
MAX_SRC_TOKENS = 80          # long sentences are cut into pieces of at most this many tokens
MAX_NEW_TOKENS = 320         # Santali needs many more tokens than Hindi for the same meaning
PDF_TRANSLATION_BATCH_SIZE = 2  # modest memory use for CPU-only classroom servers
INDIC_FONT = "Nirmala UI"    # covers Ol Chiki and Devanagari on Windows

LANGS = {
    "sat": {"name": "Santali", "script": "Bengali script (sat_Beng)", "ai": True},
    "ho": {"name": "Ho", "script": "Devanagari", "ai": False},
    "mun": {"name": "Mundari", "script": "Devanagari", "ai": False},
}

# The original adapter's corpus and evaluation report are absent, so its score
# cannot be independently reproduced or verified.
BENCHMARK = None


class ApiError(Exception):
    def __init__(self, status, msg):
        super().__init__(msg)
        self.status, self.msg = status, msg


# ---------------------------------------------------------------- text helpers
SENT_SPLIT = re.compile(r"(?<=[।!?])\s+|(?<=\.)\s+")
HAS_LETTER = re.compile(r"[\u0904-\u0939\u0958-\u0961]")        # at least one Devanagari letter
PUNCT = "।॥,.?!;:\"'()[]{}<>-–—…“”‘’/\\|"


def split_sentences(text):
    text = " ".join(str(text).split())
    return [p for p in (s.strip() for s in SENT_SPLIT.split(text)) if p]


def split_for_model(text, count_tokens, limit):
    """Cut a long sentence into pieces of at most `limit` tokens, at word boundaries.
    Without this the model would silently drop everything after its input limit."""
    pieces, cur = [], []
    for word in text.split():
        if cur and count_tokens(" ".join(cur + [word])) > limit:
            pieces.append(" ".join(cur))
            cur = [word]
        else:
            cur.append(word)
    if cur:
        pieces.append(" ".join(cur))
    return pieces


# ---------------------------------------------------------------- translation model (Santali)
class Translator:
    """Hindi -> Santali using NLLB-200 plus your LoRA adapter, with a confidence score."""

    def __init__(self, mock=False):
        self.mock = mock
        self.lock = threading.Lock()
        if mock:
            self.label = "demo mode (fake translations)"
            return
        import torch
        from transformers import AutoModelForSeq2SeqLM, AutoTokenizer

        self.torch = torch
        self.tok = AutoTokenizer.from_pretrained(BASE_MODEL, src_lang="hin_Deva")
        model = AutoModelForSeq2SeqLM.from_pretrained(BASE_MODEL)
        if ADAPTER_DIR.exists():
            # PEFT checks torch.distributed.tensor.DTensor while attaching LoRA.
            # Some CPU builds expose that namespace only after it is imported.
            import torch.distributed.tensor
            from peft import PeftModel

            model = PeftModel.from_pretrained(model, str(ADAPTER_DIR)).merge_and_unload()
            self.label = "fine-tuned"
        else:
            print("WARNING: folder 'lora_santali' not found next to app.py. Using the base model without fine-tuning.")
            self.label = "base model (not fine-tuned)"
        self.model = model.eval()
        # Language tags are placed by hand, exactly as in training.
        self.src_id = self.tok.convert_tokens_to_ids("hin_Deva")
        self.tgt_id = self.tok.convert_tokens_to_ids("sat_Beng")
        self.eos = self.tok.eos_token_id
        if self.tok.unk_token_id in (self.src_id, self.tgt_id):
            raise RuntimeError("Tokenizer does not know the hin_Deva / sat_Beng language tags.")

    def count_tokens(self, text):
        return len(self.tok(text, add_special_tokens=False)["input_ids"])

    def _translate_piece_batch(self, texts, source_id=None, target_id=None):
        torch = self.torch
        source_id = self.src_id if source_id is None else source_id
        target_id = self.tgt_id if target_id is None else target_id
        pad_id = self.tok.pad_token_id
        if pad_id is None:
            raise RuntimeError("The local tokenizer has no padding token for batched translation.")
        rows = []
        for text in texts:
            body = self.tok(text, add_special_tokens=False)["input_ids"][: MAX_SRC_TOKENS + 20]
            rows.append([source_id] + body + [self.eos])
        width = max(map(len, rows))
        input_ids = torch.full((len(rows), width), pad_id, dtype=torch.long)
        attention_mask = torch.zeros((len(rows), width), dtype=torch.long)
        for index, row in enumerate(rows):
            input_ids[index, :len(row)] = torch.tensor(row, dtype=torch.long)
            attention_mask[index, :len(row)] = 1

        with self.lock, torch.no_grad():
            out = self.model.generate(
                input_ids=input_ids,
                attention_mask=attention_mask,
                forced_bos_token_id=target_id,
                max_new_tokens=MAX_NEW_TOKENS,
                num_beams=1,
                no_repeat_ngram_size=3,
                output_scores=True,
                return_dict_in_generate=True,
                pad_token_id=pad_id,
            )
            scores = self.model.compute_transition_scores(
                out.sequences, out.scores, normalize_logits=True)

        results = []
        for index, sequence in enumerate(out.sequences):
            tokens = sequence.tolist()
            # NLLB starts the decoder with EOS; look for a generated EOS after it.
            eos_position = next((pos for pos in range(1, len(tokens)) if tokens[pos] == self.eos), None)
            cut_off = eos_position is None
            # Score generated text (and EOS), excluding the forced language tag
            # and any padding added after a shorter item in the batch finished.
            transition_scores = scores[index, 1:eos_position] if eos_position is not None else scores[index, 1:]
            conf = float(torch.exp(transition_scores.mean())) if len(transition_scores) else 0.0
            if not torch.isfinite(torch.tensor(conf)):
                conf = 0.0
            results.append((self.tok.decode(sequence, skip_special_tokens=True), conf, cut_off))
        return results

    def translate_many(self, texts, batch_size=PDF_TRANSLATION_BATCH_SIZE):
        """Translate texts in small padded batches; retain each text's lowest-piece score."""
        if self.mock:
            return [("[demo] " + text, 0.35 + (sum(map(ord, text)) % 55) / 100) for text in texts]

        pieces_by_text = [split_for_model(text, self.count_tokens, MAX_SRC_TOKENS) for text in texts]
        pieces = []
        owners = []
        for owner, text_pieces in enumerate(pieces_by_text):
            for piece in text_pieces:
                pieces.append(piece)
                owners.append(owner)
        translations = [[] for _ in texts]
        confidences = [[] for _ in texts]
        truncated = [False for _ in texts]
        size = max(1, min(int(batch_size), 8))
        for start in range(0, len(pieces), size):
            batch = self._translate_piece_batch(pieces[start:start + size])
            for owner, (translated, confidence, cut_off) in zip(owners[start:start + size], batch):
                translations[owner].append(translated)
                confidences[owner].append(confidence)
                truncated[owner] = truncated[owner] or cut_off
        return [
            (" ".join(translations[index]),
             0.0 if truncated[index] or not confidences[index] else min(confidences[index]))
            for index in range(len(texts))
        ]

    def translate_tagged_many(self, texts, source_language, target_language,
                              batch_size=PDF_TRANSLATION_BATCH_SIZE):
        """Translate via explicit NLLB tags for the English-to-Hindi bridge."""
        if self.mock:
            return ["[demo Hindi] " + str(text) for text in texts]
        source_id = self.tok.convert_tokens_to_ids(source_language)
        target_id = self.tok.convert_tokens_to_ids(target_language)
        if self.tok.unk_token_id in (source_id, target_id):
            raise RuntimeError(f"The local tokenizer does not support {source_language} -> {target_language}.")
        pieces_by_text = [split_for_model(text, self.count_tokens, MAX_SRC_TOKENS) for text in texts]
        pieces = []
        owners = []
        for owner, text_pieces in enumerate(pieces_by_text):
            for piece in text_pieces:
                pieces.append(piece)
                owners.append(owner)
        translated = [[] for _ in texts]
        size = max(1, min(int(batch_size), 8))
        for start in range(0, len(pieces), size):
            batch = self._translate_piece_batch(
                pieces[start:start + size], source_id=source_id, target_id=target_id,
            )
            for owner, (text, _confidence, _cut_off) in zip(owners[start:start + size], batch):
                translated[owner].append(text)
        return [" ".join(parts) for parts in translated]

    def translate(self, text):
        """Return (santali_text, confidence between 0 and 1). Confidence 0 means 'cut off, check closely'."""
        return self.translate_many([text], batch_size=1)[0]


_translator = None
_translator_lock = threading.Lock()


def get_translator():
    global _translator
    with _translator_lock:
        if _translator is None:
            _translator = Translator(mock=MOCK)
    return _translator


def get_translator_label():
    """Report the configured model without loading its weights into memory."""
    if MOCK:
        return "demo mode (fake translations)"
    translator = _translator
    if translator is not None:
        return translator.label
    return "fine-tuned" if ADAPTER_DIR.exists() else "base model (not fine-tuned)"


def release_translator():
    """Unload NLLB weights before the larger optional TTS model is started."""
    global _translator
    with _translator_lock:
        translator = _translator
        if translator is None:
            return False
        model = getattr(translator, "model", None)
        if model is not None:
            with translator.lock:
                translator.model = None
                translator.tok = None
                _translator = None
        else:
            _translator = None
        del model
        del translator
        gc.collect()
    return True


# ---------------------------------------------------------------- database
SCHEMA = """
CREATE TABLE IF NOT EXISTS docs (
    id      INTEGER PRIMARY KEY AUTOINCREMENT,
    name    TEXT NOT NULL,
    kind    TEXT NOT NULL,
    created REAL NOT NULL,
    units   INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS segs (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    doc_id       INTEGER NOT NULL,
    unit         INTEGER NOT NULL,
    pos          INTEGER NOT NULL,
    text         TEXT NOT NULL,
    translatable INTEGER NOT NULL
);
CREATE INDEX IF NOT EXISTS segs_doc ON segs(doc_id, unit, pos);
CREATE INDEX IF NOT EXISTS segs_text ON segs(text);
CREATE TABLE IF NOT EXISTS tr (
    seg_id      INTEGER NOT NULL,
    lang        TEXT NOT NULL,
    draft       TEXT,
    final       TEXT,
    confidence  REAL,
    model       TEXT,
    status      TEXT NOT NULL DEFAULT 'pending',
    edited      INTEGER NOT NULL DEFAULT 0,
    edit_ratio  REAL NOT NULL DEFAULT 0,
    from_memory INTEGER NOT NULL DEFAULT 0,
    note        TEXT,
    updated     REAL,
    PRIMARY KEY (seg_id, lang)
);
CREATE TABLE IF NOT EXISTS doc_langs (
    doc_id INTEGER NOT NULL,
    lang   TEXT NOT NULL,
    state  TEXT NOT NULL,
    PRIMARY KEY (doc_id, lang)
);
CREATE TABLE IF NOT EXISTS glossary (
    id     INTEGER PRIMARY KEY AUTOINCREMENT,
    lang   TEXT NOT NULL,
    hindi  TEXT NOT NULL,
    word   TEXT NOT NULL,
    source TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'pending',
    review_round INTEGER NOT NULL DEFAULT 0,
    rejection_note TEXT,
    UNIQUE (lang, hindi, word)
);
CREATE TABLE IF NOT EXISTS glossary_review_events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    glossary_id INTEGER NOT NULL,
    reviewer_role TEXT NOT NULL,
    hindi TEXT NOT NULL,
    word TEXT NOT NULL,
    review_round INTEGER NOT NULL DEFAULT 0,
    reviewed REAL NOT NULL,
    action TEXT NOT NULL DEFAULT 'approve',
    note TEXT,
    UNIQUE (glossary_id, review_round, reviewer_role)
);
"""


@contextmanager
def db():
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    try:
        yield con
        con.commit()
    finally:
        con.close()


UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
with db() as _con:
    _con.execute("PRAGMA journal_mode=WAL")
    _con.executescript(SCHEMA)
    glossary_columns = {row["name"] for row in _con.execute("PRAGMA table_info(glossary)").fetchall()}
    if "status" not in glossary_columns:
        _con.execute("ALTER TABLE glossary ADD COLUMN status TEXT NOT NULL DEFAULT 'pending'")
    if "review_round" not in glossary_columns:
        _con.execute("ALTER TABLE glossary ADD COLUMN review_round INTEGER NOT NULL DEFAULT 0")
    if "rejection_note" not in glossary_columns:
        _con.execute("ALTER TABLE glossary ADD COLUMN rejection_note TEXT")
    glossary_event_columns = {
        row["name"] for row in _con.execute("PRAGMA table_info(glossary_review_events)").fetchall()
    }
    if "action" not in glossary_event_columns:
        _con.execute("ALTER TABLE glossary_review_events ADD COLUMN action TEXT NOT NULL DEFAULT 'approve'")
    if "note" not in glossary_event_columns:
        _con.execute("ALTER TABLE glossary_review_events ADD COLUMN note TEXT")
    _con.execute("CREATE INDEX IF NOT EXISTS glossary_review_work ON glossary_review_events(glossary_id, review_round, reviewer_role)")
    _con.execute("CREATE INDEX IF NOT EXISTS glossary_status_work ON glossary(status, lang, id)")


# ---------------------------------------------------------------- reading documents
KINDS = {".docx": "docx", ".pptx": "pptx", ".pdf": "pdf", ".txt": "txt"}


class Unit:
    """One paragraph-like piece of a document, with a way to write a replacement back."""

    __slots__ = ("text", "set", "meta")

    def __init__(self, text, setter=None, meta=None):
        self.text, self.set, self.meta = text, setter, meta or {}


def _need(module, package):
    try:
        return importlib.import_module(module)
    except ImportError:
        raise ApiError(
            500,
            f"The app is missing the '{package}' package. In the terminal run: python -m pip install {package}   "
            "then stop and restart app.py.",
        )


def _set_rfonts(rpr, name):
    from docx.oxml.ns import qn

    rf = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), name)


# --- Word
def _docx_paragraphs(container):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    el = container._element
    body = el.body if hasattr(el, "body") else el
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, container)
        elif tag == "tbl":
            for row in Table(child, container).rows:
                seen = []
                for cell in row.cells:                  # merged cells repeat; read each once
                    if any(cell._tc is s for s in seen):
                        continue
                    seen.append(cell._tc)
                    yield from _docx_paragraphs(cell)


def _docx_set_text(p, new):
    from docx.oxml.ns import qn

    ts = list(p._p.iter(qn("w:t")))
    if ts:
        ts[0].text = new
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for t in ts[1:]:
            t.text = ""
        run_el = ts[0].getparent()
    else:
        run_el = p.add_run(new)._r
    _set_rfonts(run_el.get_or_add_rPr(), INDIC_FONT)


def docx_units(path):
    docx = _need("docx", "python-docx")
    doc = docx.Document(str(path))
    units = []
    for p in _docx_paragraphs(doc):
        text = " ".join(p.text.split())
        if text:
            units.append(Unit(text, lambda new, p=p: _docx_set_text(p, new)))
    return doc, units


# --- PowerPoint
def _pptx_frames(shapes):
    for sh in shapes:
        try:
            is_group = sh.shape_type == 6
        except Exception:
            is_group = False
        if is_group:
            yield from _pptx_frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _pptx_set_text(para, new):
    from lxml import etree
    from pptx.oxml.ns import qn

    runs = list(para.runs)
    if runs:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
        run = runs[0]
    else:
        run = para.add_run()
        run.text = new
    run.font.name = INDIC_FONT
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is not None and rpr.find(qn("a:cs")) is None:
        cs = etree.SubElement(rpr, qn("a:cs"))
        cs.set("typeface", INDIC_FONT)
        latin.addnext(cs)


def pptx_units(path):
    pptx = _need("pptx", "python-pptx")
    prs = pptx.Presentation(str(path))
    units = []
    for slide in prs.slides:
        for tf in _pptx_frames(slide.shapes):
            for para in tf.paragraphs:
                text = " ".join(para.text.split())
                if text:
                    units.append(Unit(text, lambda new, para=para: _pptx_set_text(para, new)))
    return prs, units


# --- PDF (only the PyMuPDF calls live in pdf_blocks, so the rest can be tested without it)
def pdf_blocks(path):
    """[(page_number, text), ...] in the order the PDF stores them."""
    try:
        import pymupdf as fitz
    except ImportError:
        fitz = _need("fitz", "pymupdf")
    out = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            for b in page.get_text("blocks"):
                if b[6] == 0:                                   # 0 = text block
                    t = " ".join(b[4].split())
                    if t:
                        out.append((page.number + 1, t))
    return out


def pdf_units(path):
    blocks = pdf_blocks(path)
    if sum(len(t) for _, t in blocks) < 50:
        raise ApiError(
            400,
            "This legacy PDF reader only extracts selectable text. For scanned pages, run the supported FastAPI/PWA "
            "server with optional local Tesseract OCR, or upload a text-based PDF.",
        )
    return [Unit(t, None, {"page": pg}) for pg, t in blocks]


# --- Plain text: one unit per line
def txt_units(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ApiError(400, "This text file is not UTF-8. Open it in Notepad and use Save As with encoding UTF-8.")
    lines = text.splitlines()
    units = []
    for i, line in enumerate(lines):
        t = " ".join(line.split())
        if t:
            units.append(Unit(t, lambda new, i=i: lines.__setitem__(i, new)))
    return lines, units


def extract_units(kind, path):
    try:
        if kind == "docx":
            return docx_units(path)[1]
        if kind == "pptx":
            return pptx_units(path)[1]
        if kind == "pdf":
            return pdf_units(path)
        return txt_units(Path(path).read_bytes())[1]
    except ApiError:
        raise
    except Exception as e:
        raise ApiError(400, f"This file could not be read as .{kind}. It may be damaged, or not a real .{kind} file "
                            f"({type(e).__name__}).")


def doc_path(doc_id, kind):
    return UPLOAD_DIR / str(doc_id) / f"original.{kind}"


# ---------------------------------------------------------------- documents
def create_document(filename, data):
    if not data:
        raise ApiError(400, "The file is empty.")
    name = Path(filename or "").name.strip() or "document"
    ext = Path(name).suffix.lower()
    if ext == ".doc":
        raise ApiError(400, "Old .doc files are not supported. Open it in Word and save it as .docx.")
    kind = KINDS.get(ext)
    if not kind:
        raise ApiError(400, "Upload a .docx, .pptx, .pdf or .txt file.")
    with db() as con:
        doc_id = con.execute("INSERT INTO docs (name, kind, created) VALUES (?,?,?)", (name, kind, time.time())).lastrowid
    path = doc_path(doc_id, kind)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    try:
        units = extract_units(kind, path)
        rows = []
        for ui, u in enumerate(units):
            for pi, s in enumerate(split_sentences(u.text)):
                rows.append((doc_id, ui, pi, s, int(bool(HAS_LETTER.search(s)))))
        if not any(r[4] for r in rows):
            raise ApiError(400, "No Hindi text was found in this file. This version translates Hindi (Devanagari) documents.")
        if len(rows) > MAX_SENTENCES:
            raise ApiError(400, f"This document has more than {MAX_SENTENCES} sentences. Split it into smaller files.")
        with db() as con:
            con.executemany("INSERT INTO segs (doc_id, unit, pos, text, translatable) VALUES (?,?,?,?,?)", rows)
            con.execute("UPDATE docs SET units=? WHERE id=?", (len(units), doc_id))
    except Exception:
        delete_doc(doc_id)
        raise
    return get_doc(doc_id)


def create_pasted(text, title):
    text = str(text or "")
    if not text.strip():
        raise ApiError(400, "Paste some Hindi text first.")
    title = re.sub(r"[^\w\- ]", "", str(title or "").strip(), flags=re.UNICODE)[:60].strip() or "Pasted text"
    return create_document(title + ".txt", text.encode("utf-8"))


def delete_doc(doc_id):
    with db() as con:
        con.execute("DELETE FROM tr WHERE seg_id IN (SELECT id FROM segs WHERE doc_id=?)", (doc_id,))
        con.execute("DELETE FROM segs WHERE doc_id=?", (doc_id,))
        con.execute("DELETE FROM doc_langs WHERE doc_id=?", (doc_id,))
        con.execute("DELETE FROM docs WHERE id=?", (doc_id,))
    shutil.rmtree(UPLOAD_DIR / str(doc_id), ignore_errors=True)
    return {"deleted": doc_id}


def lang_counts(con, doc_id, lang):
    r = con.execute(
        """SELECT COUNT(*) AS total,
                  COALESCE(SUM(t.draft IS NOT NULL AND t.draft != ''), 0) AS drafted,
                  COALESCE(SUM(t.status = 'approved'), 0) AS approved,
                  COALESCE(SUM(t.draft IS NULL), 0) AS waiting,
                  COALESCE(SUM(t.status = 'pending' AND t.draft IS NOT NULL AND t.draft != ''), 0) AS pending_drafts,
                  COALESCE(SUM(t.from_memory), 0) AS from_memory
           FROM tr t JOIN segs s ON s.id = t.seg_id WHERE s.doc_id=? AND t.lang=?""",
        (doc_id, lang),
    ).fetchone()
    return dict(r)


def _doc_summary(con, d):
    sentences = con.execute("SELECT COUNT(*) FROM segs WHERE doc_id=?", (d["id"],)).fetchone()[0]
    translatable = con.execute("SELECT COUNT(*) FROM segs WHERE doc_id=? AND translatable=1", (d["id"],)).fetchone()[0]
    langs = {}
    for r in con.execute("SELECT lang, state FROM doc_langs WHERE doc_id=?", (d["id"],)).fetchall():
        langs[r["lang"]] = {"state": r["state"], **lang_counts(con, d["id"], r["lang"])}
    return {"id": d["id"], "name": d["name"], "kind": d["kind"], "created": d["created"], "units": d["units"],
            "sentences": sentences, "translatable": translatable, "langs": langs}


def get_doc(doc_id):
    with db() as con:
        d = con.execute("SELECT * FROM docs WHERE id=?", (doc_id,)).fetchone()
        if not d:
            raise ApiError(404, "That document no longer exists. Go back to the document list.")
        return _doc_summary(con, d)


def list_docs():
    with db() as con:
        return [_doc_summary(con, d) for d in con.execute("SELECT * FROM docs ORDER BY id DESC").fetchall()]


# ---------------------------------------------------------------- translating
TM_SQL = """SELECT t.final FROM tr t JOIN segs s ON s.id = t.seg_id
            WHERE s.text=? AND t.lang=? AND t.status='approved' AND t.final IS NOT NULL AND t.final != ''
            ORDER BY t.updated DESC LIMIT 1"""

JOBS = queue.Queue()


def start_language(doc_id, lang):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    ai = LANGS[lang]["ai"]
    now = time.time()
    with db() as con:
        if not con.execute("SELECT 1 FROM docs WHERE id=?", (doc_id,)).fetchone():
            raise ApiError(404, "That document no longer exists.")
        segs = con.execute("SELECT id, text FROM segs WHERE doc_id=? AND translatable=1 ORDER BY id", (doc_id,)).fetchall()
        for s in segs:
            if con.execute("SELECT 1 FROM tr WHERE seg_id=? AND lang=?", (s["id"], lang)).fetchone():
                continue
            memory = con.execute(TM_SQL, (s["text"], lang)).fetchone()
            if memory:        # a teacher already approved this exact sentence: reuse it
                con.execute(
                    "INSERT INTO tr (seg_id, lang, draft, final, confidence, model, status, from_memory, updated) "
                    "VALUES (?,?,?,?,1.0,'memory','approved',1,?)",
                    (s["id"], lang, memory["final"], memory["final"], now),
                )
            else:
                con.execute("INSERT INTO tr (seg_id, lang, draft, status, updated) VALUES (?,?,?, 'pending', ?)",
                            (s["id"], lang, None if ai else "", now))
        waiting = con.execute(
            "SELECT COUNT(*) FROM tr t JOIN segs s ON s.id=t.seg_id WHERE s.doc_id=? AND t.lang=? AND t.draft IS NULL",
            (doc_id, lang),
        ).fetchone()[0]
        state = "queued" if (ai and waiting) else ("done" if ai else "ready")
        con.execute("INSERT OR REPLACE INTO doc_langs (doc_id, lang, state) VALUES (?,?,?)", (doc_id, lang, state))
    if state == "queued":
        JOBS.put((doc_id, lang))
    return get_doc(doc_id)


def _set_state(doc_id, lang, state):
    with db() as con:
        con.execute("UPDATE doc_langs SET state=? WHERE doc_id=? AND lang=?", (state, doc_id, lang))


def run_job(doc_id, lang):
    _set_state(doc_id, lang, "running")
    translator = get_translator()
    while True:
        with db() as con:
            row = con.execute(
                "SELECT s.id, s.text FROM segs s JOIN tr t ON t.seg_id=s.id "
                "WHERE s.doc_id=? AND t.lang=? AND t.draft IS NULL ORDER BY s.id LIMIT 1",
                (doc_id, lang),
            ).fetchone()
        if not row:
            break
        try:
            out, conf = translator.translate(row["text"])
            note = None
        except Exception as e:           # one bad sentence must not stop the whole document
            out, conf, note = "", 0.0, f"The model failed on this sentence ({type(e).__name__}). Type the translation yourself."
        with db() as con:
            con.execute(
                "UPDATE tr SET draft=?, confidence=?, model=?, note=?, updated=? WHERE seg_id=? AND lang=?",
                (out, conf, translator.label, note, time.time(), row["id"], lang),
            )
    _set_state(doc_id, lang, "done")


def worker_loop():
    while True:
        doc_id, lang = JOBS.get()
        try:
            run_job(doc_id, lang)
        except Exception:
            traceback.print_exc()
            _set_state(doc_id, lang, "error")


def requeue_unfinished():
    with db() as con:
        rows = con.execute("SELECT doc_id, lang FROM doc_langs WHERE state IN ('queued','running')").fetchall()
    for r in rows:
        if LANGS.get(r["lang"], {}).get("ai"):
            JOBS.put((r["doc_id"], r["lang"]))


# ---------------------------------------------------------------- review
def _glossary_map(con, lang):
    m = {}
    for r in con.execute("SELECT hindi, word FROM glossary WHERE lang=?", (lang,)).fetchall():
        m.setdefault(r["hindi"], []).append(r["word"])
    return m


def get_work(doc_id, lang, filt="pending", limit=30):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    limit = max(1, min(int(limit), 100))
    with db() as con:
        if not con.execute("SELECT 1 FROM docs WHERE id=?", (doc_id,)).fetchone():
            raise ApiError(404, "That document no longer exists.")
        counts = lang_counts(con, doc_id, lang)
        where = "s.doc_id=? AND t.lang=? AND t.draft IS NOT NULL"
        order = "s.id ASC"
        if filt == "pending":
            where += " AND t.status='pending'"
            if LANGS[lang]["ai"]:
                order = "t.confidence ASC, s.id ASC"       # least sure first
        rows = con.execute(
            "SELECT s.id AS seg_id, s.unit, s.text, t.draft, t.final, t.confidence, t.status, t.model, "
            "t.from_memory, t.note FROM segs s JOIN tr t ON t.seg_id = s.id "
            f"WHERE {where} ORDER BY {order} LIMIT ?",
            (doc_id, lang, limit),
        ).fetchall()
        gloss = {} if LANGS[lang]["ai"] else _glossary_map(con, lang)
    items = []
    for r in rows:
        it = dict(r)
        hits = []
        if gloss:
            seen = set()
            for tok in it["text"].split():
                tok = tok.strip(PUNCT)
                if tok in gloss and tok not in seen:
                    seen.add(tok)
                    hits += [{"hindi": tok, "word": w} for w in gloss[tok]]
        it["hits"] = hits
        items.append(it)
    return {"counts": counts, "items": items}


def approve_segment(seg_id, lang, final):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    final = " ".join(str(final or "").split())
    if not final:
        raise ApiError(400, "The translation is empty. Type a translation before approving.")
    with db() as con:
        row = con.execute("SELECT * FROM tr WHERE seg_id=? AND lang=?", (seg_id, lang)).fetchone()
        if not row:
            raise ApiError(404, "That sentence is not in this language's list. Reload the page.")
        draft = (row["draft"] or "").strip()
        edited = int(bool(draft) and final != draft)
        ratio = (1 - difflib.SequenceMatcher(None, draft, final).ratio()) if edited else 0.0
        model = row["model"] if draft else "human"
        memory = row["from_memory"] if final == draft else 0
        con.execute(
            "UPDATE tr SET final=?, status='approved', edited=?, edit_ratio=?, model=?, from_memory=?, updated=? "
            "WHERE seg_id=? AND lang=?",
            (final, edited, ratio, model, memory, time.time(), seg_id, lang),
        )
    return {"seg_id": seg_id, "edited": bool(edited), "edit_ratio": ratio}


def approve_all(doc_id, lang):
    if lang not in LANGS or not LANGS[lang]["ai"]:
        raise ApiError(400, "Approve all is only available for languages with a machine draft.")
    with db() as con:
        cur = con.execute(
            "UPDATE tr SET final=draft, status='approved', edited=0, updated=? "
            "WHERE lang=? AND status='pending' AND draft IS NOT NULL AND draft != '' "
            "AND seg_id IN (SELECT id FROM segs WHERE doc_id=?)",
            (time.time(), lang, doc_id),
        )
        return {"approved": cur.rowcount}


# ---------------------------------------------------------------- student view, corpus, stats
def get_approved(lang):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    with db() as con:
        rows = con.execute(
            "SELECT d.id AS doc_id, d.name, s.text AS hindi, t.final AS target FROM tr t "
            "JOIN segs s ON s.id=t.seg_id JOIN docs d ON d.id=s.doc_id "
            "WHERE t.lang=? AND t.status='approved' ORDER BY d.id, s.unit, s.pos",
            (lang,),
        ).fetchall()
    docs, order = {}, []
    for r in rows:
        if r["doc_id"] not in docs:
            docs[r["doc_id"]] = {"doc_id": r["doc_id"], "name": r["name"], "items": []}
            order.append(r["doc_id"])
        docs[r["doc_id"]]["items"].append({"hindi": r["hindi"], "target": r["target"]})
    return [docs[i] for i in order]


def corpus_jsonl(lang):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    with db() as con:
        rows = con.execute(
            "SELECT d.name, s.text, t.final, t.edited, t.model FROM tr t JOIN segs s ON s.id=t.seg_id "
            "JOIN docs d ON d.id=s.doc_id WHERE t.lang=? AND t.status='approved' ORDER BY d.id, s.unit, s.pos",
            (lang,),
        ).fetchall()
    lines = [
        json.dumps({"hindi": r["text"], "target": r["final"], "language": lang, "teacher_edited": bool(r["edited"]),
                    "source": r["model"], "document": r["name"]}, ensure_ascii=False)
        for r in rows
    ]
    return ("\n".join(lines) + "\n") if lines else ""


def get_stats():
    out = {"languages": {}, "documents": 0, "benchmark": BENCHMARK, "model": get_translator_label()}
    with db() as con:
        one = lambda sql, *a: con.execute(sql, a).fetchone()[0]
        out["documents"] = one("SELECT COUNT(*) FROM docs")
        for lang in LANGS:
            q = "FROM tr WHERE lang=?"
            out["languages"][lang] = {
                "drafted": one(f"SELECT COUNT(*) {q} AND draft IS NOT NULL AND draft != ''", lang),
                "approved": one(f"SELECT COUNT(*) {q} AND status='approved'", lang),
                "edited": one(f"SELECT COUNT(*) {q} AND status='approved' AND edited=1", lang),
                "human_written": one(f"SELECT COUNT(*) {q} AND status='approved' AND model='human'", lang),
                "from_memory": one(f"SELECT COUNT(*) {q} AND from_memory=1", lang),
                "avg_conf": one(f"SELECT AVG(confidence) {q} AND model != 'memory' AND model != 'human'", lang),
                "conf_edited": one(f"SELECT AVG(confidence) {q} AND status='approved' AND edited=1", lang),
                "conf_kept": one(f"SELECT AVG(confidence) {q} AND status='approved' AND edited=0 "
                                 "AND from_memory=0 AND model != 'human'", lang),
                "glossary": one("SELECT COUNT(*) FROM glossary WHERE lang=?", lang),
            }
    return out


# ---------------------------------------------------------------- word lists
def seed_glossary():
    with db() as con:
        for filename, lang in (("ho_hindi_pairs_raw.csv", "ho"), ("mundari_pairs_raw.csv", "mun")):
            path = VOCAB_DIR / filename
            if not path.exists():
                continue
            with open(path, encoding="utf-8", newline="") as f:
                for row in csv.DictReader(f):
                    word, hindi = (row.get("target") or "").strip(), (row.get("hindi") or "").strip()
                    if len(hindi) < 2 or (len(word) < 2 and len(hindi) <= 3):
                        continue
                    if re.search(r"[A-Za-z0-9\u0966-\u096F\u0964]", word + hindi) or " " in word or " " in hindi:
                        continue
                    con.execute("INSERT OR IGNORE INTO glossary (lang, hindi, word, source) VALUES (?,?,?,?)",
                                (lang, hindi, word, "NCERT primer (read by OCR)"))


seed_glossary()


def add_glossary(lang, hindi, word, source="added by a reviewer"):
    if lang not in LANGS or LANGS[lang]["ai"]:
        raise ApiError(400, "Choose Ho or Mundari.")
    hindi, word = " ".join(str(hindi or "").split()), " ".join(str(word or "").split())
    if not hindi or not word:
        raise ApiError(400, "Fill in both the Hindi word and the word in the other language.")
    if " " in hindi:
        raise ApiError(400, "The Hindi side must be a single word.")
    if len(word) > 60 or len(hindi) > 60:
        raise ApiError(400, "That entry is too long for a word list.")
    with db() as con:
        cur = con.execute("INSERT OR IGNORE INTO glossary (lang, hindi, word, source) VALUES (?,?,?,?)",
                          (lang, hindi, word, source))
    return {"added": cur.rowcount == 1}


def search_glossary(q, lang):
    q = (q or "").strip()
    with db() as con:
        counts = {r["lang"]: r["n"] for r in con.execute(
            "SELECT lang, COUNT(*) AS n FROM glossary WHERE status='approved' GROUP BY lang")}
        results = []
        if q:
            sql, args = "SELECT id,lang, hindi, word, source FROM glossary WHERE status='approved' AND (hindi LIKE ? OR word LIKE ?)", [f"%{q}%"] * 2
            if lang in LANGS:
                sql, args = sql + " AND lang=?", args + [lang]
            results = [dict(r) for r in con.execute(sql + " ORDER BY lang, hindi LIMIT 40", args).fetchall()]
    return {"counts": counts, "results": results}


def list_approved_glossary(limit=5000):
    with db() as con:
        counts = {r["lang"]: r["n"] for r in con.execute(
            "SELECT lang, COUNT(*) AS n FROM glossary WHERE status='approved' GROUP BY lang")}
        rows = con.execute(
            "SELECT id,lang,hindi,word,source FROM glossary WHERE status='approved' ORDER BY lang,hindi LIMIT ?",
            (max(1, min(int(limit), 10000)),),
        ).fetchall()
    total = sum(counts.values())
    return {"counts": counts, "results": [dict(row) for row in rows], "complete": len(rows) >= total}


# ---------------------------------------------------------------- exporting
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "txt": "text/plain; charset=utf-8",
}


def build_unit_texts(con, doc_id, lang, mode):
    rows = con.execute(
        "SELECT s.unit, s.text, s.translatable, t.draft, t.final, t.status FROM segs s "
        "LEFT JOIN tr t ON t.seg_id = s.id AND t.lang=? WHERE s.doc_id=? ORDER BY s.unit, s.pos",
        (lang, doc_id),
    ).fetchall()
    units, done, total = {}, 0, 0
    for r in rows:
        piece = r["text"]
        if r["translatable"]:
            total += 1
            if r["status"] == "approved" and r["final"]:
                piece, done = r["final"], done + 1
            elif mode == "drafts" and r["draft"]:
                piece, done = r["draft"], done + 1
        units.setdefault(r["unit"], []).append(piece)
    return {u: " ".join(ps) for u, ps in units.items()}, done, total


def _pdf_as_docx(units, texts):
    docx = _need("docx", "python-docx")
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name, style.font.size = INDIC_FONT, Pt(12)
    _set_rfonts(style.element.get_or_add_rPr(), INDIC_FONT)
    last_page = None
    for i, u in enumerate(units):
        page = u.meta.get("page")
        if page != last_page:
            note = d.add_paragraph()
            r = note.add_run(f"Page {page}")
            r.italic = True
            r.font.size = Pt(9)
            last_page = page
        p = d.add_paragraph()
        r = p.add_run(texts.get(i, u.text))
        _set_rfonts(r._r.get_or_add_rPr(), INDIC_FONT)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def export_document(doc_id, lang, mode):
    if lang not in LANGS:
        raise ApiError(400, "Unknown language.")
    mode = "drafts" if mode == "drafts" else "approved"
    with db() as con:
        d = con.execute("SELECT * FROM docs WHERE id=?", (doc_id,)).fetchone()
        if not d:
            raise ApiError(404, "That document no longer exists.")
        texts, done, total = build_unit_texts(con, doc_id, lang, mode)
    if done == 0:
        raise ApiError(400, "Nothing is approved in this language yet. Approve some sentences first.")
    kind, path = d["kind"], doc_path(doc_id, d["kind"])
    if not path.exists():
        raise ApiError(404, "The original file is missing from the uploads folder.")
    stem = Path(d["name"]).stem
    if kind == "docx":
        doc, units = docx_units(path)
        for i, u in enumerate(units):
            if i in texts and texts[i] != u.text:
                u.set(texts[i])
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), f"{stem}.{lang}.docx", MIME["docx"]
    if kind == "pptx":
        prs, units = pptx_units(path)
        for i, u in enumerate(units):
            if i in texts and texts[i] != u.text:
                u.set(texts[i])
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue(), f"{stem}.{lang}.pptx", MIME["pptx"]
    if kind == "pdf":
        return _pdf_as_docx(pdf_units(path), texts), f"{stem}.{lang}.docx", MIME["docx"]
    lines, units = txt_units(path.read_bytes())
    for i, u in enumerate(units):
        if i in texts and texts[i] != u.text:
            u.set(texts[i])
    return ("\n".join(lines) + "\n").encode("utf-8"), f"{stem}.{lang}.txt", MIME["txt"]


# ---------------------------------------------------------------- web server
SW_JS = """
const CACHE = 'bhasha-setu-v2';
self.addEventListener('install', e => e.waitUntil(
  caches.open(CACHE).then(c => c.add('/')).then(() => self.skipWaiting())));
self.addEventListener('activate', e => e.waitUntil(self.clients.claim()));
self.addEventListener('fetch', e => {
  const u = new URL(e.request.url);
  if (e.request.method !== 'GET' || u.pathname !== '/') return;
  e.respondWith(fetch(e.request)
    .then(r => { const copy = r.clone(); caches.open(CACHE).then(c => c.put('/', copy)); return r; })
    .catch(() => caches.match('/')));
});
"""


def health():
    libs = {"python-docx": "docx", "python-pptx": "pptx", "pymupdf": "pymupdf"}
    missing = [pkg for pkg, mod in libs.items()
               if importlib.util.find_spec(mod) is None and not (mod == "pymupdf" and importlib.util.find_spec("fitz"))]
    return {"model": get_translator_label(), "langs": LANGS, "missing": missing}


class Handler(BaseHTTPRequestHandler):
    server_version = "BhashaSetu/2.0"

    def log_message(self, *args):
        pass

    def _send(self, status, body, ctype, headers=None):
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, status=200):
        self._send(status, json.dumps(obj, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8")

    def _file(self, data, filename, ctype):
        ascii_name = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
        self._send(200, data, ctype,
                   {"Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"})

    def _raw(self, limit):
        n = int(self.headers.get("Content-Length") or 0)
        if n > limit:
            raise ApiError(413, f"That is too large. The limit is {limit // 1_000_000} MB.")
        return self.rfile.read(n) if n else b""

    def _json_body(self):
        raw = self._raw(3_000_000)
        try:
            data = json.loads(raw.decode("utf-8") or "{}")
        except ValueError:
            raise ApiError(400, "The request body is not valid JSON.")
        if not isinstance(data, dict):
            raise ApiError(400, "The request body must be a JSON object.")
        return data

    def _route(self, method):
        url = urlparse(self.path)
        path, qs = url.path, parse_qs(url.query)
        arg = lambda k, default="": (qs.get(k) or [default])[0]
        try:
            if method == "GET":
                if path == "/":
                    return self._send(200, (HERE / "index.html").read_bytes(), "text/html; charset=utf-8")
                if path == "/sw.js":
                    return self._send(200, SW_JS.encode("utf-8"), "application/javascript; charset=utf-8")
                if path == "/favicon.ico":
                    return self._send(204, b"", "image/x-icon")
                if path == "/api/health":
                    return self._json(health())
                if path == "/api/docs":
                    return self._json(list_docs())
                if path == "/api/stats":
                    return self._json(get_stats())
                if path == "/api/approved":
                    return self._json(get_approved(arg("lang", "sat")))
                if path == "/api/glossary":
                    return self._json(search_glossary(arg("q"), arg("lang")))
                if path == "/api/corpus":
                    lang = arg("lang", "sat")
                    return self._file(corpus_jsonl(lang).encode("utf-8"), f"corpus_{lang}.jsonl", "application/x-ndjson")
                m = re.fullmatch(r"/api/doc/(\d+)", path)
                if m:
                    return self._json(get_doc(int(m.group(1))))
                m = re.fullmatch(r"/api/doc/(\d+)/work", path)
                if m:
                    return self._json(get_work(int(m.group(1)), arg("lang", "sat"), arg("filter", "pending"), arg("limit", "30") or 30))
                m = re.fullmatch(r"/api/export/(\d+)/(\w+)", path)
                if m:
                    return self._file(*export_document(int(m.group(1)), m.group(2), arg("mode", "approved")))
            elif method == "POST":
                if path == "/api/upload":
                    return self._json(create_document(unquote(self.headers.get("X-Filename", "")), self._raw(MAX_UPLOAD)))
                body = self._json_body()
                if path == "/api/paste":
                    return self._json(create_pasted(body.get("text"), body.get("title")))
                if path == "/api/glossary":
                    return self._json(add_glossary(body.get("lang"), body.get("hindi"), body.get("word")))
                m = re.fullmatch(r"/api/doc/(\d+)/(start|approve_all|delete)", path)
                if m:
                    doc_id, action = int(m.group(1)), m.group(2)
                    if action == "start":
                        return self._json(start_language(doc_id, body.get("lang")))
                    if action == "approve_all":
                        return self._json(approve_all(doc_id, body.get("lang")))
                    return self._json(delete_doc(doc_id))
                m = re.fullmatch(r"/api/seg/(\d+)/(\w+)/approve", path)
                if m:
                    return self._json(approve_segment(int(m.group(1)), m.group(2), body.get("final")))
            raise ApiError(404, "Not found.")
        except ApiError as e:
            self._json({"detail": e.msg}, e.status)
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            traceback.print_exc()
            self._json({"detail": "Something went wrong on the server. The terminal window shows the details."}, 500)

    def do_GET(self):
        self._route("GET")

    def do_POST(self):
        self._route("POST")


def selftest():
    tr = get_translator()
    long_sentence = ("बच्चे सुबह उठकर नहाते हैं, खाना खाते हैं, किताबें अपने बस्ते में रखते हैं, और फिर अपने दोस्तों के साथ "
                     "हँसते-खेलते हुए गाँव की कच्ची सड़क से होकर स्कूल जाते हैं, जहाँ उनके शिक्षक उन्हें रोज़ पढ़ाते हैं।")
    for sentence in ["बच्चे स्कूल जाते हैं", "मेरा नाम राम है", long_sentence]:
        t0 = time.time()
        santali, conf = tr.translate(sentence)
        pieces = "" if MOCK else f"  ({len(split_for_model(sentence, tr.count_tokens, MAX_SRC_TOKENS))} piece(s))"
        print(f"{sentence}\n  -> {santali}\n  confidence {conf:.2f}   {time.time() - t0:.1f}s{pieces}")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    if "--selftest" in sys.argv:
        print("Starting in demo mode (no model)." if MOCK else
              "Loading the model. The first run downloads about 2.4 GB, so this can take a few minutes...")
        tr = get_translator()
        print("Model ready:", tr.label)
        selftest()
        sys.exit(0)
    raise SystemExit("Use 'python api.py' to run the supported Bhasha Setu PWA and FastAPI server.")
